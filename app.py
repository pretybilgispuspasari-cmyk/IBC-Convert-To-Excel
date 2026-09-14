import io
import os
import re
from flask import Flask, render_template, request, send_file
import pdfplumber
import pandas as pd
from pdf2image import convert_from_bytes
import pytesseract
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl import load_workbook

base_dir = os.path.dirname(os.path.abspath(__file__))
template_dir = os.path.join(base_dir, 'templates')

app = Flask(__name__, template_folder=template_dir)

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/convert', methods=['POST'])
def convert_pdf_to_excel():
    if 'file' not in request.files:
        return "No file uploaded", 400
    
    file = request.files['file']
    selected_bank = request.form.get('bank', 'umum').lower()

    try:
        pdf_bytes = file.read()
        full_text_collected = ""

        with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
            for page in pdf.pages:
                extracted = page.extract_text(layout=True)
                if extracted:
                    full_text_collected += extracted + "\n"

        if not full_text_collected.strip():
            images = convert_from_bytes(pdf_bytes)
            for img in images:
                ocr_text = pytesseract.image_to_string(img)
                full_text_collected += ocr_text + "\n"

        if not full_text_collected.strip():
            return "Tabel atau teks tidak ditemukan di dalam PDF.", 400

        lines = [line.strip() for line in full_text_collected.split('\n') if line.strip()]
        parsed_rows = []

        def parse_num(val_str):
            try:
                clean = val_str.replace(',', '').replace(' ', '')
                return float(clean)
            except:
                return ""

        # --- LOGIKA KHUSUS BANK BNI ---
        if selected_bank == 'bni':
            # Contoh format tanggal BNI: "01 Jan 2022" atau "31 Jan 2022"
            bni_date_pattern = re.compile(r'^(\d{2}\s+[A-Za-z]{3}\s+\d{4})')
            amount_pattern = re.compile(r'([\d,\.]+\.\d{2}|\d+)')

            current_date = ""
            current_uraian = ""
            current_debit = ""
            current_credit = ""
            current_saldo = ""

            def save_bni_row():
                nonlocal current_date, current_uraian, current_debit, current_credit, current_saldo
                if current_date or current_uraian:
                    parsed_rows.append([
                        current_date,
                        current_uraian,
                        '', # Keterangan kosong atau gabung
                        current_debit,
                        current_credit,
                        current_saldo
                    ])
                current_date = ""
                current_uraian = ""
                current_debit = ""
                current_credit = ""
                current_saldo = ""

            for line in lines:
                skip_words = ['TANGGAL', 'URAIAN', 'DEBET', 'KREDIT', 'SALDO', 'HALAMAN', 'REKENING']
                if any(w in line.upper() for w in skip_words) and not re.match(r'^\d{2}\s+[A-Za-z]{3}\s+\d{4}', line):
                    continue

                date_match = bni_date_pattern.match(line)
                if date_match:
                    save_bni_row()
                    current_date = date_match.group(1)
                    rem_text = line[len(current_date):].strip()

                    # Cari angka-angka di ujung baris (biasanya Saldo, Debet/Kredit)
                    # Di BNI urutannya: [Tanggal] [Uraian/Keterangan] [Debet?] [Kredit?] [Saldo]
                    # Atau SALDO AWAL hanya punya saldo di ujung.
                    tokens = rem_text.split()
                    # Ambil angka dari belakang ke depan
                    nums_found = []
                    text_parts = []
                    
                    for tok in reversed(tokens):
                        # Cek apakah token berupa angka (bisa ada koma/titik)
                        cleaned_tok = tok.replace(',', '')
                        if cleaned_tok.replace('.', '', 1).isdigit() and len(nums_found) < 3:
                            nums_found.insert(0, parse_num(tok))
                        else:
                            text_parts.insert(0, tok)

                    current_uraian = " ".join(text_parts)

                    if current_uraian.upper() == 'SALDO AWAL':
                        if nums_found:
                            current_saldo = nums_found[-1]
                    else:
                        # Analisis posisi angka (Debet vs Kredit vs Saldo)
                        if len(nums_found) == 3:
                            current_debit = nums_found[0] if nums_found[0] != 0 else ""
                            current_credit = nums_found[1] if nums_found[1] != 0 else ""
                            current_saldo = nums_found[2]
                        elif len(nums_found) == 2:
                            # Bisa jadi debet + saldo atau kredit + saldo
                            # Cek kata kunci di uraian
                            if any(k in current_uraian.upper() for k in ['CH', 'INT CR', 'CHEQUE', 'CR']):
                                current_credit = nums_found[0]
                            else:
                                current_debit = nums_found[0]
                            current_saldo = nums_found[1]
                        elif len(nums_found) == 1:
                            current_saldo = nums_found[0]
                else:
                    # Baris lanjutan uraian jika ada
                    if current_uraian:
                        current_uraian += " " + line

            save_bni_row()

        # --- LOGIKA BANK BCA / UMUM (SEBELUMNYA) ---
        else:
            current_date = ""
            current_uraian = ""
            current_desc_list = []
            current_debit = ""
            current_credit = ""
            current_saldo = ""

            date_pattern = re.compile(r'^(\d{2}/\d{2}(?:/\d{4}?)?)')
            amount_pattern = re.compile(r'([\d,\.]+\.\d{2})')

            known_uraians = [
                'SALDO AWAL', 'KR OTOMATIS', 'TRSF E-BANKING DB', 'TRSF E-BANKING CR',
                'BI-FAST CR', 'BI-FAST DB', 'BIAYA ADM', 'BIAYA ADMIN', 'ATM DB', 'ATM CR'
            ]

            skip_keywords = [
                'HALAMAN', 'PERIODE', 'TANGGAL', 'KETERANGAN', 'BCA berhak', 
                'CABANG', 'KCU', 'NO. REKENING', 'MATA UANG', 'REKENING', 
                'KEC ', 'RT ', 'DSN ', 'BLITAR', 'INDONESIA', 'CATATAN:'
            ]

            def save_current_row():
                nonlocal current_date, current_uraian, current_desc_list, current_debit, current_credit, current_saldo
                if current_date or current_uraian or current_desc_list:
                    parsed_rows.append([
                        current_date,
                        current_uraian,
                        " ".join(current_desc_list),
                        current_debit,
                        current_credit,
                        current_saldo
                    ])
                current_date = ""
                current_uraian = ""
                current_desc_list = []
                current_debit = ""
                current_credit = ""
                current_saldo = ""

            for line in lines:
                if any(kw in line.upper() for kw in skip_keywords) and not ('TANGGAL :' in line.upper()):
                    continue

                date_match = date_pattern.match(line)
                if date_match:
                    save_current_row()
                    current_date = date_match.group(1)
                    rem_text = line[len(current_date):].strip()

                    found_uraian = ""
                    for u in known_uraians:
                        if rem_text.upper().startswith(u):
                            found_uraian = u
                            rem_text = rem_text[len(u):].strip()
                            break
                    
                    if not found_uraian:
                        parts = rem_text.split()
                        if parts:
                            found_uraian = parts[0]
                            rem_text = " ".join(parts[1:])

                    current_uraian = found_uraian

                    amounts = amount_pattern.findall(rem_text)
                    desc_text = rem_text
                    for amt in amounts:
                        desc_text = desc_text.replace(amt, "").strip()
                    
                    if desc_text:
                        current_desc_list.append(desc_text)
                    
                    if current_uraian.upper() == 'SALDO AWAL':
                        if amounts:
                            current_saldo = parse_num(amounts[-1])
                    else:
                        if "DB" in line.upper():
                            if len(amounts) >= 1:
                                current_debit = parse_num(amounts[-1])
                        elif "CR" in line.upper() or len(amounts) >= 2:
                            if len(amounts) >= 1:
                                current_credit = parse_num(amounts[0])
                        if len(amounts) > 1:
                            current_saldo = parse_num(amounts[-1])
                else:
                    if 'TANGGAL :' in line.upper() or any(u in line.upper() for u in known_uraians):
                        current_desc_list.append(line)
                    else:
                        if not line.startswith("•"):
                            current_desc_list.append(line)

            save_current_row()

        if not parsed_rows:
            parsed_rows = [['Tidak ada data transaksi yang dapat diparsing', '', '', '', '', '']]

        headers = ['Tanggal', 'Uraian', 'Keterangan', 'Debit', 'Kredit', 'Saldo']
        df = pd.DataFrame(parsed_rows, columns=headers)

        excel_stream = io.BytesIO()
        with pd.ExcelWriter(excel_stream, engine='openpyxl') as writer:
            df.to_excel(writer, index=False, sheet_name='Mutasi')
        
        # --- STYLING OPENPYXL ---
        excel_stream.seek(0)
        wb = load_workbook(excel_stream)
        ws = wb['Mutasi']

        blue_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        white_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        regular_font = Font(name="Calibri", size=11)
        
        thin_border = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )

        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = blue_fill
            cell.font = white_font
            cell.alignment = Alignment(horizontal='center', vertical='center')

        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=len(headers)):
            for cell in row:
                cell.font = regular_font
                cell.border = thin_border
                if cell.column in [4, 5, 6]:
                    cell.alignment = Alignment(horizontal='right', vertical='center')
                    if isinstance(cell.value, (int, float)):
                        cell.number_format = '#,##0.00'
                elif cell.column == 1:
                    cell.alignment = Alignment(horizontal='center', vertical='center')
                else:
                    cell.alignment = Alignment(horizontal='left', vertical='center')

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

        final_stream = io.BytesIO()
        wb.save(final_stream)
        final_stream.seek(0)

        return send_file(
            final_stream,
            download_name=f"konversi_{selected_bank}.xlsx",
            as_attachment=True,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
    except Exception as e:
        return f"Terjadi kesalahan: {str(e)}", 500

import os

if __name__ == '__main__':
    port = int(os.environ.get('PORT', 5000))
    app.run(host='0.0.0.0', port=port)